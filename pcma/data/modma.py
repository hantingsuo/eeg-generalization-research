"""MODMA EGI raw loader and prepared-feature cache utilities.

The Lanzhou MODMA ERP files are EGI NetStation ``.raw`` recordings with 129 EEG
channels plus event channels. The v1 feature table uses fixed cue-locked windows
from the three dot-probe cue channels (``fcue``, ``hcue``, ``scue``).
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

import numpy as np
from openpyxl import load_workbook
from scipy.signal import welch


DEFAULT_MODMA_ROOT = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "854301_EEG_128Channels_ERP_Lanzhou_2015"
    / "EEG_128channels_ERP_lanzhou_2015"
)

DEFAULT_BANDS: tuple[tuple[str, float, float], ...] = (
    ("delta", 1.0, 4.0),
    ("theta", 4.0, 8.0),
    ("alpha", 8.0, 13.0),
    ("beta", 13.0, 30.0),
    ("gamma", 30.0, 45.0),
)

DEFAULT_EVENT_CHANNELS: tuple[str, ...] = ("fcue", "hcue", "scue")
SCALE_COLUMNS: tuple[str, ...] = ("PHQ-9", "CTQ-SF", "LES", "SSRS", "GAD-7", "PSQI")


@dataclass(frozen=True)
class ModmaSubjectRecord:
    subject_id: str
    population: str
    raw_path: Path | None
    age: float | None = None
    gender: str | None = None
    education_years: float | None = None
    scales: dict[str, float] = field(default_factory=dict)


@dataclass(frozen=True)
class ModmaFeatureConfig:
    event_channels: tuple[str, ...] = DEFAULT_EVENT_CHANNELS
    tmin: float = 0.0
    tmax: float = 1.0
    max_events_per_label: int = 30
    bands: tuple[tuple[str, float, float], ...] = DEFAULT_BANDS


@dataclass
class ModmaFeatureData:
    F: np.ndarray
    y_emotion: np.ndarray | None
    subject: np.ndarray
    population: np.ndarray
    phq9_by_subject: dict[str, float]
    event: np.ndarray | None = None
    scales_by_subject: dict[str, dict[str, float]] = field(default_factory=dict)
    feature_names: list[str] | None = None
    source: str | None = None
    config: dict[str, object] | None = None


def _subject_id_from_name(path: str | Path) -> str:
    name = Path(path).name
    m = re.match(r"^(0\d{7})", name)
    if not m:
        raise ValueError(f"cannot parse MODMA subject id from {name!r}")
    return m.group(1)


def _subject_id(value: object) -> str:
    if value is None:
        raise ValueError("missing subject id")
    text = str(value).strip()
    m = re.search(r"0\d{7}", text)
    if m:
        return m.group(0)
    if text.isdigit():
        return f"{int(text):08d}"
    raise ValueError(f"cannot normalize subject id {value!r}")


def _as_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return out if np.isfinite(out) else None


def _header_key(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    low = text.lower()
    if low == "subject id":
        return "subject_id"
    if low == "type":
        return "type"
    if low == "age":
        return "age"
    if low == "gender":
        return "gender"
    if low.startswith("education"):
        return "education_years"
    for scale in SCALE_COLUMNS:
        if low == scale.lower():
            return scale
    return None


def discover_modma_files(root: str | Path = DEFAULT_MODMA_ROOT) -> tuple[dict[str, Path], Path | None]:
    """Return raw files keyed by subject ID and the subject-info xlsx path."""
    root = Path(root)
    raw_paths = sorted(root.glob("*.raw"))
    raw_by_subject = {_subject_id_from_name(p): p for p in raw_paths}
    xlsx = sorted(root.glob("*subject*.xlsx"))
    if not xlsx:
        xlsx = sorted(root.glob("*.xlsx"))
    return raw_by_subject, (xlsx[0] if xlsx else None)


def load_modma_metadata(root: str | Path = DEFAULT_MODMA_ROOT) -> list[ModmaSubjectRecord]:
    """Load the subject-information workbook and join rows to raw files."""
    raw_by_subject, xlsx = discover_modma_files(root)
    if xlsx is None:
        raise FileNotFoundError(f"no MODMA subject-information xlsx found under {root}")
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        raw_header = next(rows)
    except StopIteration as exc:
        raise ValueError(f"empty MODMA workbook: {xlsx}") from exc
    header = [_header_key(v) for v in raw_header]
    idx = {key: i for i, key in enumerate(header) if key is not None}
    required = {"subject_id", "type", "PHQ-9"}
    missing = required - set(idx)
    if missing:
        raise ValueError(f"MODMA workbook missing columns: {sorted(missing)}")

    records: list[ModmaSubjectRecord] = []
    for row in rows:
        if idx["subject_id"] >= len(row) or row[idx["subject_id"]] is None:
            continue
        sid = _subject_id(row[idx["subject_id"]])
        population = str(row[idx["type"]]).strip().upper()
        if population not in {"HC", "MDD"}:
            raise ValueError(f"{sid}: unexpected MODMA type {population!r}")
        scales = {}
        for scale in SCALE_COLUMNS:
            if scale in idx:
                value = _as_float(row[idx[scale]] if idx[scale] < len(row) else None)
                if value is not None:
                    scales[scale] = value
        records.append(
            ModmaSubjectRecord(
                subject_id=sid,
                population=population,
                raw_path=raw_by_subject.get(sid),
                age=_as_float(row[idx["age"]]) if "age" in idx and idx["age"] < len(row) else None,
                gender=str(row[idx["gender"]]).strip() if "gender" in idx and idx["gender"] < len(row) and row[idx["gender"]] is not None else None,
                education_years=_as_float(row[idx["education_years"]])
                if "education_years" in idx and idx["education_years"] < len(row)
                else None,
                scales=scales,
            )
        )
    return records


def _rising_edges(x: np.ndarray) -> np.ndarray:
    active = np.asarray(x) > 0
    return np.flatnonzero(np.diff(np.r_[False, active.astype(bool)].astype(np.int8)) == 1)


def egi_event_samples(raw, event_channels: Iterable[str] = DEFAULT_EVENT_CHANNELS) -> dict[str, np.ndarray]:
    """Extract rising-edge sample indices for named EGI event channels."""
    names = list(event_channels)
    missing = [ch for ch in names if ch not in raw.ch_names]
    if missing:
        raise ValueError(f"raw file missing MODMA event channels: {missing}")
    picks = [raw.ch_names.index(ch) for ch in names]
    stim = raw.get_data(picks=picks)
    return {ch: _rising_edges(stim[i]) for i, ch in enumerate(names)}


def _balanced_event_subset(samples: np.ndarray, max_events: int) -> np.ndarray:
    samples = np.asarray(samples, dtype=int)
    if max_events <= 0 or len(samples) <= max_events:
        return samples
    idx = np.linspace(0, len(samples) - 1, max_events)
    return samples[np.unique(np.rint(idx).astype(int))]


def bandpower_feature_names(ch_names: list[str], bands: tuple[tuple[str, float, float], ...] = DEFAULT_BANDS) -> list[str]:
    names = []
    for prefix in ("de", "rel_power"):
        for ch in ch_names:
            for band, _, _ in bands:
                names.append(f"{prefix}_{band}_{ch}")
    return names


def bandpower_de_features(epoch: np.ndarray, sfreq: float, bands: tuple[tuple[str, float, float], ...] = DEFAULT_BANDS) -> np.ndarray:
    """Return per-channel band DE and relative power for one epoch."""
    x = np.asarray(epoch, dtype=float)
    if x.ndim != 2:
        raise ValueError(f"epoch must be 2D (n_channels, n_times), got {x.shape}")
    x = x - x.mean(axis=1, keepdims=True)
    nperseg = min(int(round(sfreq)), x.shape[1])
    freqs, psd = welch(x, fs=sfreq, nperseg=nperseg, axis=1)
    powers = []
    for _, lo, hi in bands:
        mask = (freqs >= lo) & (freqs < hi)
        if not np.any(mask):
            powers.append(np.full(x.shape[0], 1e-12))
        else:
            powers.append(np.trapz(psd[:, mask], freqs[mask], axis=1) + 1e-12)
    power = np.stack(powers, axis=1)
    de = 0.5 * np.log(2.0 * np.pi * np.e * power)
    rel = power / (power.sum(axis=1, keepdims=True) + 1e-12)
    return np.concatenate([de.reshape(-1), rel.reshape(-1)]).astype(np.float32)


def _feature_config_dict(config: ModmaFeatureConfig) -> dict[str, object]:
    return {
        "event_channels": list(config.event_channels),
        "tmin": config.tmin,
        "tmax": config.tmax,
        "max_events_per_label": config.max_events_per_label,
        "bands": [[name, lo, hi] for name, lo, hi in config.bands],
    }


def load_modma_raw_features(
    root: str | Path = DEFAULT_MODMA_ROOT,
    config: ModmaFeatureConfig | None = None,
    progress: Callable[[str], None] | None = None,
) -> ModmaFeatureData:
    """Build cue-locked MODMA bandpower features from local EGI raw files."""
    import mne

    config = config or ModmaFeatureConfig()
    records = load_modma_metadata(root)
    rows: list[np.ndarray] = []
    labels: list[str] = []
    subjects: list[str] = []
    populations: list[str] = []
    events: list[str] = []
    feature_names: list[str] | None = None
    scales_by_subject = {r.subject_id: dict(r.scales) for r in records}
    phq9_by_subject = {r.subject_id: r.scales["PHQ-9"] for r in records if "PHQ-9" in r.scales}

    for i, record in enumerate(records, start=1):
        if record.raw_path is None:
            raise FileNotFoundError(f"{record.subject_id}: no joined MODMA raw file")
        if progress is not None:
            progress(f"[{i}/{len(records)}] reading {record.subject_id} {record.population}")
        raw = mne.io.read_raw_egi(record.raw_path, preload=False, verbose="ERROR")
        sfreq = float(raw.info["sfreq"])
        eeg_picks = mne.pick_types(raw.info, eeg=True, stim=False, exclude=[])
        ch_names = [raw.ch_names[p] for p in eeg_picks]
        if feature_names is None:
            feature_names = bandpower_feature_names(ch_names, config.bands)
        elif feature_names != bandpower_feature_names(ch_names, config.bands):
            raise ValueError(f"{record.subject_id}: EEG channel list differs from prior subjects")

        event_samples = egi_event_samples(raw, config.event_channels)
        start_offset = int(round(config.tmin * sfreq))
        stop_offset = int(round(config.tmax * sfreq))
        if stop_offset <= start_offset:
            raise ValueError("feature window must have positive duration")
        eeg = raw.get_data(picks=eeg_picks)
        for event_name in config.event_channels:
            selected = _balanced_event_subset(event_samples[event_name], config.max_events_per_label)
            for sample in selected:
                start = int(sample) + start_offset
                stop = int(sample) + stop_offset
                if start < 0 or stop > eeg.shape[1]:
                    continue
                feat = bandpower_de_features(eeg[:, start:stop], sfreq=sfreq, bands=config.bands)
                rows.append(feat)
                labels.append(event_name)
                subjects.append(record.subject_id)
                populations.append(record.population)
                events.append(event_name)

    if not rows:
        raise ValueError("no MODMA feature rows were extracted")
    return ModmaFeatureData(
        F=np.vstack(rows).astype(np.float32),
        y_emotion=np.asarray(labels),
        subject=np.asarray(subjects),
        population=np.asarray(populations),
        phq9_by_subject=phq9_by_subject,
        event=np.asarray(events),
        scales_by_subject=scales_by_subject,
        feature_names=feature_names,
        source=str(Path(root)),
        config=_feature_config_dict(config),
    )


def save_modma_feature_cache(path: str | Path, data: ModmaFeatureData) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    scale_subjects = np.asarray(sorted(data.scales_by_subject), dtype=object)
    scale_names = np.asarray(SCALE_COLUMNS, dtype=object)
    scale_values = np.full((len(scale_subjects), len(scale_names)), np.nan, dtype=float)
    for i, sid in enumerate(scale_subjects):
        for j, scale in enumerate(scale_names):
            value = data.scales_by_subject.get(str(sid), {}).get(str(scale))
            if value is not None:
                scale_values[i, j] = float(value)
    phq9_subject = np.asarray(sorted(data.phq9_by_subject), dtype=object)
    phq9 = np.asarray([data.phq9_by_subject[str(sid)] for sid in phq9_subject], dtype=float)
    payload = {
        "F": np.asarray(data.F, dtype=np.float32),
        "subject": np.asarray(data.subject).astype(str),
        "population": np.asarray(data.population).astype(str),
        "phq9_subject": phq9_subject,
        "phq9": phq9,
        "scale_subject": scale_subjects,
        "scale_names": scale_names,
        "scale_values": scale_values,
        "feature_names": np.asarray(data.feature_names or [], dtype=object),
        "source": np.asarray([data.source or ""], dtype=object),
    }
    if data.config is not None:
        payload["config_json"] = np.asarray([json.dumps(data.config, ensure_ascii=False)], dtype=object)
    if data.y_emotion is not None:
        payload["y_emotion"] = np.asarray(data.y_emotion).astype(str)
    if data.event is not None:
        payload["event"] = np.asarray(data.event).astype(str)
    np.savez_compressed(path, **payload)


def load_modma_feature_cache(path: str | Path) -> ModmaFeatureData:
    """Load prepared MODMA features from an NPZ cache."""
    d = np.load(path, allow_pickle=True)
    required = ["F", "subject", "population", "phq9_subject", "phq9"]
    missing = [key for key in required if key not in d]
    if missing:
        raise ValueError(f"missing required MODMA feature-cache keys: {missing}")
    F = np.asarray(d["F"], dtype=float)
    y = np.asarray(d["y_emotion"]).astype(str) if "y_emotion" in d else None
    subject = np.asarray(d["subject"]).astype(str)
    population = np.asarray(d["population"]).astype(str)
    if F.ndim != 2:
        raise ValueError(f"F must be 2D, got {F.shape}")
    n = F.shape[0]
    for name, arr in {"subject": subject, "population": population}.items():
        if len(arr) != n:
            raise ValueError(f"{name} length {len(arr)} != n_trials {n}")
    if y is not None and len(y) != n:
        raise ValueError(f"y_emotion length {len(y)} != n_trials {n}")
    phq_subject = np.asarray(d["phq9_subject"]).astype(str)
    phq9 = np.asarray(d["phq9"], dtype=float)
    if len(phq_subject) != len(phq9):
        raise ValueError("phq9_subject and phq9 lengths differ")
    phq = {str(s): float(v) for s, v in zip(phq_subject, phq9) if np.isfinite(v)}
    event = np.asarray(d["event"]).astype(str) if "event" in d else None
    if event is not None and len(event) != n:
        raise ValueError(f"event length {len(event)} != n_trials {n}")

    scales_by_subject: dict[str, dict[str, float]] = {}
    if {"scale_subject", "scale_names", "scale_values"}.issubset(set(d.files)):
        scale_subjects = np.asarray(d["scale_subject"]).astype(str)
        scale_names = np.asarray(d["scale_names"]).astype(str)
        scale_values = np.asarray(d["scale_values"], dtype=float)
        for i, sid in enumerate(scale_subjects):
            scales_by_subject[str(sid)] = {
                str(name): float(scale_values[i, j])
                for j, name in enumerate(scale_names)
                if i < scale_values.shape[0] and j < scale_values.shape[1] and np.isfinite(scale_values[i, j])
            }
    else:
        scales_by_subject = {sid: {"PHQ-9": value} for sid, value in phq.items()}
    feature_names = np.asarray(d["feature_names"]).astype(str).tolist() if "feature_names" in d else None
    source = str(np.asarray(d["source"]).astype(str)[0]) if "source" in d else str(path)
    config = None
    if "config_json" in d:
        config_text = str(np.asarray(d["config_json"]).astype(str)[0])
        if config_text:
            config = json.loads(config_text)
    return ModmaFeatureData(
        F=F,
        y_emotion=y,
        subject=subject,
        population=population,
        phq9_by_subject=phq,
        event=event,
        scales_by_subject=scales_by_subject,
        feature_names=feature_names,
        source=source,
        config=config,
    )
